"""다이나믹셀 v1.x BOM (부품마다 무엇인지·수량·재질·질량·파일).
    .venv\\Scripts\\python.exe cad/redesign/tools/make_bom.py cad/redesign/dynamixel/v1.0
출력: <버전>/BOM_dynamixel_<ver>.csv, .xlsx
수량 = 원본 MJCF 배치 개수 (manifest.json). 출력 부품 질량 = 새 솔리드 부피 × PETG 1.27 g/cm³ (꽉 찬 기준, 실제 출력은 채움률만큼 가벼움).
"""
import sys, json, csv
from pathlib import Path
ROOT = Path(__file__).resolve().parents[3]
V = Path(sys.argv[1]); V = V if V.is_absolute() else ROOT / V
VER = V.name
T = Path(__file__).parent
M = json.loads((T / 'manifest.json').read_text(encoding='utf-8'))
ALIAS = {'seeed_bearing__configuration__22x16x4': 'bearing_16x22x4', 'seeed_bearing__configuration_default': 'bearing_10x15x3'}
BODY_KO = {'trunk_base': '몸통', 'yaw2roll': '왼쪽 고관절 yaw', 'hip_l': '왼쪽 고관절 roll', 'upper_leg_left': '왼쪽 허벅지', 'leg': '왼쪽 정강이',
           'ankle_left': '왼발', 'neck': '목', 'neck_pitch': '머리 pitch', 'yaw_roll_motion': '머리 yaw', 'jaw_soft': '머리 roll(머리 전체)',
           'bearing_roll': '오른쪽 고관절 yaw', 'hip_l_2': '오른쪽 고관절 roll', 'upper_leg_right': '오른쪽 허벅지', 'leg_2': '오른쪽 정강이', 'ankle_right': '오른발'}
# 이름: (분류, 무엇인지, 재질·규격, 비고)
P, S, PL, B = '출력(경질)', '출력(연질)', '판재', '구매품'
INFO = {
    'trunk_base': (PL, '몸통 바닥판. 좌우 hip_yaw XL330을 받치고 Ø19 구멍으로 혼·베어링이 지나감', '알루미늄 5052-H32 1.0 mm', '레이저/워터젯. DXF·PDF 도면'),
    'bearing_roll': (PL, '고관절 롤 베어링판 (고관절 앞면 수직판). 좌우 같은 부품', '알루미늄 5052-H32 1.0 mm', '원본 위치결정 핀 생략'),
    'upper_leg_rigidity_plate': (PL, '허벅지 보강판 (허벅지 바깥 옆면, hip_pitch·knee 서보를 잇는 측판). 좌우 같은 부품', '알루미늄 5052-H32 1.0 mm', 'v1.0: 무릎 접힘 여유 홈'),
    'neck': (PL, '목 옆판 (목 서보 2개를 잇는 판, 양옆 2장)', '알루미늄 5052-H32 2.0 mm', '테두리 1.5t로 외형 확대 24.3×13.3'),
    'right_shell': (P, '몸통 오른쪽 외피(셸)', 'PETG 또는 PLA+', ''),
    'left_shell': (P, '몸통 왼쪽 외피(셸)', 'PETG 또는 PLA+', ''),
    'power_support': (P, '몸통 전원부 지지대 (배터리·스위치·퓨즈 자리)', 'PETG 또는 PLA+', ''),
    'banana_pcb_locker': (P, '몸통 안 보드/배선 고정 막대 (원본 명칭 banana PCB locker, 정확한 용도 미확인)', 'PETG 또는 PLA+', ''),
    'yaw2roll': (P, '고관절 yaw→roll 연결 브래킷 (hip_yaw 혼에 붙고 hip_roll 서보를 잡음). 좌우 공용', 'PETG 또는 PLA+', ''),
    'hip_l': (P, '고관절 roll→pitch 브래킷 (Ø16 축이 베어링에 끼움). 좌우 공용', 'PETG 또는 PLA+', 'v1.0: Ø16 축 정확한 원통'),
    'upper_leg_left': (P, '왼쪽 허벅지 링크 (hip_pitch·knee XL330 홀더)', 'PETG 또는 PLA+', ''),
    'upper_leg_right': (P, '오른쪽 허벅지 링크 (hip_pitch·knee XL330 홀더)', 'PETG 또는 PLA+', ''),
    'leg': (P, '정강이 링크 (knee 혼 ↔ ankle XL330). 좌우 공용', 'PETG 또는 PLA+', ''),
    'ankle_left': (P, '왼쪽 발목 브래킷 (ankle 혼 ↔ 발판)', 'PETG 또는 PLA+', ''),
    'ankle_right': (P, '오른쪽 발목 브래킷', 'PETG 또는 PLA+', ''),
    'foot_left': (P, '왼쪽 발판', 'PETG 또는 PLA+', ''),
    'foot_right': (P, '오른쪽 발판', 'PETG 또는 PLA+', ''),
    'sole_left': (S, '왼쪽 발바닥 (미끄럼 방지)', 'TPU 95A', ''),
    'sole_right': (S, '오른쪽 발바닥', 'TPU 95A', ''),
    'neck_pitch': (P, '목 위 브래킷 (head_pitch 서보·혼, Ø16 축 베어링)', 'PETG 또는 PLA+', 'v1.0: Ø16·Ø18 원통 정확'),
    'yaw_roll_motion': (P, '머리 yaw–roll 브래킷 (head_yaw 혼 ↔ head_roll 서보)', 'PETG 또는 PLA+', '재설계 때 롤 혼 볼트용 Ø4 공구 구멍 추가 예정'),
    'face_part': (P, '얼굴판 (보드·카메라 장착면)', 'PETG 또는 PLA+', ''),
    'motor_support': (P, '머리 안 서보 지지대 (M2.5 ×4 체결)', 'PETG 또는 PLA+', ''),
    'noenoeil': (P, '얼굴의 눈 링 (원본 명칭 noenoeil)', 'PETG 또는 PLA+', ''),
    'm12_lens_holder': (P, 'M12 렌즈 홀더 (IMX219 카메라용)', 'PETG 또는 PLA+', ''),
    'top_head_shell': (P, '머리 윗 셸', 'PETG 또는 PLA+', ''),
    'bottom_head_shell': (P, '머리 아래 셸', 'PETG 또는 PLA+', ''),
    'jaw': (P, '부리 아래(경질)', 'PETG 또는 PLA+', ''),
    'jaw_soft': (S, '부리 아래 연질부', 'TPU 95A', ''),
    'soft_mouth_top': (S, '부리 위 연질부', 'TPU 95A', ''),
    'xl330': (B, '서보 ROBOTIS DYNAMIXEL XL330-M288-T (20×34×26 mm, 18 g)', 'ROBOTIS XL330-M288-T', '혼·탭 구멍은 실측 모델. 동봉 볼트 사용'),
    'bearing_16x22x4': (B, '볼베어링 내경16·외경22·폭4 (관절 지지)', '16×22×4 mm 볼베어링', ''),
    'bearing_10x15x3': (B, '볼베어링 내경10·외경15·폭3 (발목·머리)', '10×15×3 mm 볼베어링', ''),
    'np_f970': (B, '배터리 (모델 외형 NP-F970, 실물 NP-F550 2S)', 'Sony NP-F 호환 7.4 V', '외형만 모델링'),
    'pcb__raspberry_pi_zero_2_w': (B, '메인 보드 자리 (65×30 폼팩터: Radxa Zero 3W 또는 Pi Zero 2W)', 'Radxa Zero 3W', '외형·장착 구멍만'),
    'elec_rpi_robot_hat_pcb': (B, '로봇 HAT PCB (반이중 버스·전원·오디오, 직접 설계)', '65×30 자작 PCB', '외형·장착 구멍만 (hardware/hat)'),
    'lens': (B, 'M12 광각 렌즈 (IMX219 카메라)', 'M12 렌즈', '외형만'),
    'speaker': (B, '소형 스피커 35×25×7 mm', '35×25×7 스피커', '외형만'),
}
SCREWS = [('M2×6 탭볼트', 83, 'XL330 동봉품 사용 (혼 PHS M2×6 TAP)'), ('M2×8 탭볼트', 18, 'XL330 동봉품 사용 (프레임 PHS M2×8 TAP)'),
          ('M2×5', 10, ''), ('M2×4', 4, '별도 구매'), ('M2.5×6', 11, '별도 구매')]
count, bodies = {}, {}
for b in M['bodies']:
    for g in b['geoms']:
        n = ALIAS.get(g['mesh'], g['mesh'])
        count[n] = count.get(n, 0) + 1
        bodies.setdefault(n, [])
        if BODY_KO[b['name']] not in bodies[n]:
            bodies[n].append(BODY_KO[b['name']])
rows = []
order = [k for k in INFO if INFO[k][0] == PL] + [k for k in INFO if INFO[k][0] in (P, S)] + [k for k in INFO if INFO[k][0] == B]
for i, n in enumerate(order, 1):
    cat, what, mat, note = INFO[n]
    vol = mass = None
    vj = T / 'build' / 'verify' / f'{n}.json'
    if cat in (P, S) and vj.exists():
        vol = json.loads(vj.read_text())['vol_new'] / 1000
        mass = vol * (1.27 if cat == P else 1.21)
    lj = T / 'build' / f'{n}_loops.json'
    if cat == PL and lj.exists():
        d = json.loads(lj.read_text(encoding='utf-8')); mass = d['mass_g']; vol = d['volume_mm3'] / 1000
    pd = V / 'parts' / n
    files = ', '.join(sorted(f.name for f in pd.glob('*') if f.suffix in ('.step', '.stl', '.dxf', '.pdf'))) if pd.exists() else ''
    rows.append({'번호': i, '부품명': n, '무엇인지': what, '분류': cat, '수량': count.get(n, 0), '재질·규격': mat,
                 '들어가는 곳': ', '.join(bodies.get(n, [])), '부피 cm³': round(vol, 2) if vol else '', '질량 g/개': round(mass, 2) if mass else ('18' if n == 'xl330' else ''),
                 '파일 (parts/<부품명>/)': files, '비고': note})
for nm, q, note in SCREWS:
    rows.append({'번호': len(rows) + 1, '부품명': nm, '무엇인지': '체결 나사', '분류': '체결품', '수량': q, '재질·규격': nm, '들어가는 곳': '조립 가이드 단계별',
                 '부피 cm³': '', '질량 g/개': '', '파일 (parts/<부품명>/)': '', '비고': note + ' (docs/assembly 계산값)'})
hdr = list(rows[0].keys())
with open(V / f'BOM_dynamixel_{VER}.csv', 'w', encoding='utf-8-sig', newline='') as f:
    w = csv.DictWriter(f, hdr); w.writeheader(); w.writerows(rows)
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
wb = Workbook(); ws = wb.active; ws.title = f'BOM {VER}'
ws.append([f'Microduck 다이나믹셀 {VER} BOM — 어셈블리 1_microduck_dynamixel_{VER} (Fusion 개인/microduck_dynamixel_{VER})'])
ws['A1'].font = Font(bold=True, size=13)
ws.append(hdr)
for c in ws[2]:
    c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='305496'); c.alignment = Alignment(wrap_text=True, vertical='center')
fill = {PL: 'DDEBF7', P: 'E2EFDA', S: 'FFF2CC', B: 'FCE4D6', '체결품': 'EDEDED'}
for r in rows:
    ws.append([r[h] for h in hdr])
    for c in ws[ws.max_row]:
        c.fill = PatternFill('solid', fgColor=fill[r['분류']]); c.alignment = Alignment(wrap_text=True, vertical='top')
for col, wdt in zip('ABCDEFGHIJK', (5, 26, 52, 10, 6, 24, 22, 9, 9, 34, 30)):
    ws.column_dimensions[col].width = wdt
ws.freeze_panes = 'C3'
tot = {}
for r in rows:
    if isinstance(r['질량 g/개'], (int, float)):
        tot[r['분류']] = tot.get(r['분류'], 0) + r['질량 g/개'] * r['수량']
ws.append([]); ws.append(['', '분류별 질량 합계 (CAD, 출력물은 꽉 찬 기준)'])
for k, v in tot.items():
    ws.append(['', k, round(v, 1)])
ws.append(['', '서보 XL330 ×15', 18 * 15])
wb.save(V / f'BOM_dynamixel_{VER}.xlsx')
print('rows', len(rows), 'mass by category', {k: round(v, 1) for k, v in tot.items()}, '+ XL330', 18 * count.get('xl330', 0))
